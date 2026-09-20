import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    BG_COLOR = RGBColor(8, 10, 16)
    CARD_BG = RGBColor(16, 20, 30)
    CARD_BORDER = RGBColor(38, 45, 65)
    IQOO_AMBER = RGBColor(247, 181, 0)
    IQOO_ORANGE = RGBColor(255, 62, 0)
    IQOO_CYAN = RGBColor(0, 229, 255)
    IQOO_GREEN = RGBColor(0, 230, 118)
    WHITE = RGBColor(255, 255, 255)
    GRAY_TEXT = RGBColor(165, 175, 195)
    ACCENT_RED = RGBColor(255, 23, 68)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR
        return bg

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    def add_header(slide, title_text, category_text="iQOO EDGETUTOR • ON-DEVICE EDUCATION AI"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = IQOO_AMBER

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.7))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 1: THE HOOK (ZERO INTERNET AI TUTOR)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.7), Inches(2.2), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = IQOO_AMBER
    bar.line.fill.background()

    tbox = s1.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.7), Inches(2.2))
    tf = tbox.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "⚡ iQOO EdgeTutor"
    p0.font.size = Pt(44)
    p0.font.bold = True
    p0.font.color.rgb = WHITE

    p1 = tf.add_paragraph()
    p1.text = "What If Your Smartphone Could Tutor You with Zero Internet?"
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = IQOO_AMBER
    p1.space_before = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "An On-Device AI Mentor for CBSE 9-10 & JEE Problem Solving in 100% Airplane Mode"
    p2.font.size = Pt(14)
    p2.font.color.rgb = GRAY_TEXT
    p2.space_before = Pt(6)

    # 3 Pipeline Step Cards
    c1 = add_card(s1, Inches(0.8), Inches(3.4), Inches(3.6), Inches(2.5))
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "📷 1. Offline Camera Scan"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = IQOO_ORANGE
    pb = tf1.add_paragraph()
    pb.text = "• Point camera at printed textbook problem\n• Local Google ML Kit text extraction\n• Zero internet permission (0 KB Air-Gapped)"
    pb.font.size = Pt(12)
    pb.font.color.rgb = WHITE
    pb.space_before = Pt(8)

    c2 = add_card(s1, Inches(4.8), Inches(3.4), Inches(3.6), Inches(2.5))
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "🧠 2. On-Device Reasoning"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = IQOO_CYAN
    pb = tf2.add_paragraph()
    pb.text = "• Embedded NCERT curriculum vault\n• Small Language Model (SLM) reasoning\n• Grounded in CBSE standards (No hallucinations)"
    pb.font.size = Pt(12)
    pb.font.color.rgb = WHITE
    pb.space_before = Pt(8)

    c3 = add_card(s1, Inches(8.8), Inches(3.4), Inches(3.6), Inches(2.5))
    tf3 = c3.text_frame
    tf3.word_wrap = True
    p = tf3.paragraphs[0]
    p.text = "💬 3. Socratic Mentorship"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = IQOO_GREEN
    pb = tf3.add_paragraph()
    pb.text = "• Teaches derivation step-by-step\n• Asks probing questions, doesn't dump answers\n• KaTeX formatted math symbols & formulas"
    pb.font.size = Pt(12)
    pb.font.color.rgb = WHITE
    pb.space_before = Pt(8)

    fbox = s1.shapes.add_textbox(Inches(0.8), Inches(6.25), Inches(11.7), Inches(0.8))
    tff = fbox.text_frame
    pf = tff.paragraphs[0]
    pf.text = "👤 Presenter: VARANASI SAI VINAY   |   🔗 GitHub: github.com/vinayvaranasi0711-alt/iqoo-edgetutor"
    pf.font.size = Pt(12)
    pf.font.bold = True
    pf.font.color.rgb = GRAY_TEXT

    # ==========================================
    # SLIDE 2: THE REAL PROBLEM
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "The Real Problem: Why Online EdTech Fails Students", "01 / PROBLEM STATEMENT")

    prob_cards = [
        ("📶 1. The Connectivity Drop", "Cloud-dependent AI fails in college hostels, remote towns, or daily train commutes. When students study late at night and connection drops, learning halts completely.", IQOO_ORANGE),
        ("📱 2. The Distraction Trap", "Opening a browser to check a doubt exposes students to social notifications (Instagram, WhatsApp, YouTube). Deep study requires focused isolation, but cloud AI demands active connectivity.", ACCENT_RED),
        ("📋 3. Answer-Dumping vs. True Learning", "Current cloud tools act as homework copy-pasters. They dump final answers instantly without guiding the student through the fundamental physical derivations and problem-solving steps.", IQOO_AMBER)
    ]

    for i, (title, desc, col) in enumerate(prob_cards):
        top_pos = Inches(1.7 + i * 1.7)
        c = add_card(s2, Inches(0.8), top_pos, Inches(11.7), Inches(1.45))
        tfc = c.text_frame
        tfc.word_wrap = True
        pt = tfc.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(16)
        pt.font.bold = True
        pt.font.color.rgb = col
        pd = tfc.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(12)
        pd.font.color.rgb = WHITE
        pd.space_before = Pt(4)

    # ==========================================
    # SLIDE 3: HOW IT WORKS (ARCHITECTURE)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "How It Works: End-to-End On-Device Pipeline", "02 / SYSTEM ARCHITECTURE")

    arch_steps = [
        ("Step 1: Vision", "CameraX (Airplane Mode)\nOn-device Optical Scanner\nZero network permission", IQOO_ORANGE),
        ("Step 2: Local OCR", "Google ML Kit (Offline)\nExtracts formulas & text\nRuns fully in memory", IQOO_AMBER),
        ("Step 3: Curriculum RAG", "NCERT Knowledge Vault\nCBSE 9-10 & JEE 11-12\nPrevents hallucinations", RGBColor(168, 85, 247)),
        ("Step 4: SLM Engine", "Quantized Local Model\nGemma 2B / Phi-3.5 target\nSocratic step generation", IQOO_CYAN),
        ("Step 5: Socratic UI", "Interactive Student Turn\nKaTeX math typesetting\nGuided understanding", IQOO_GREEN)
    ]

    for i, (title, desc, col) in enumerate(arch_steps):
        left_pos = Inches(0.8 + i * 2.4)
        c = add_card(s3, left_pos, Inches(1.8), Inches(2.25), Inches(4.8))
        tfc = c.text_frame
        tfc.word_wrap = True
        pt = tfc.paragraphs[0]
        pt.text = title
        pt.font.size = Pt(15)
        pt.font.bold = True
        pt.font.color.rgb = col
        pd = tfc.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(11)
        pd.font.color.rgb = WHITE
        pd.space_before = Pt(12)

    # ==========================================
    # SLIDE 4: WHY iQOO & VALIDATION MATRIX
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Hardware Synergy & Implementation Validation Matrix", "03 / HARDWARE & VALIDATION")

    # Left Column: 4 Hardware Pillars
    c_hw = add_card(s4, Inches(0.8), Inches(1.8), Inches(5.3), Inches(5.0))
    tf_hw = c_hw.text_frame
    tf_hw.word_wrap = True
    pt = tf_hw.paragraphs[0]
    pt.text = "⚡ Why iQOO Hardware Fits Edge AI"
    pt.font.size = Pt(16)
    pt.font.bold = True
    pt.font.color.rgb = IQOO_AMBER

    hw_points = (
        "1. High-Bandwidth Memory (LPDDR5X up to 77 GB/s):\n"
        "   Overcomes the memory-bandwidth wall for responsive autoregressive token generation.\n\n"
        "2. Qualcomm Hexagon NPU Target:\n"
        "   Specialized INT4 quantized matrix multiplication offload at sub-watt power efficiency.\n\n"
        "3. 6000mm² Vapor Chamber Cooling:\n"
        "   Provides thermal headroom to prevent throttling during long study sessions.\n\n"
        "4. Airplane Mode Resilience:\n"
        "   Proves gaming performance can power serious student productivity."
    )
    pd = tf_hw.add_paragraph()
    pd.text = hw_points
    pd.font.size = Pt(11)
    pd.font.color.rgb = WHITE
    pd.space_before = Pt(8)

    # Right Column: Validation Matrix Table
    c_tbl = add_card(s4, Inches(6.5), Inches(1.8), Inches(6.0), Inches(5.0))
    tf_tbl = c_tbl.text_frame
    tf_tbl.word_wrap = True
    pt = tf_tbl.paragraphs[0]
    pt.text = "📊 Implementation Status (Honest Separation)"
    pt.font.size = Pt(16)
    pt.font.bold = True
    pt.font.color.rgb = IQOO_CYAN

    table_shape = s4.shapes.add_table(6, 3, Inches(6.65), Inches(2.5), Inches(5.7), Inches(4.1))
    tbl = table_shape.table

    matrix_data = [
        ["Subsystem", "Demonstrated (Now)", "Target Production"],
        ["Network", "0 KB (No Internet Manifest)", "0 KB (Strict Air-Gapped)"],
        ["Vision / OCR", "ML Kit (Offline Local)", "CameraX + ML Kit"],
        ["Curriculum", "CBSE 9-10 & JEE Vault", "Full K-12 Vector Store"],
        ["SLM Runtime", "Socratic Pipeline Emulator", "Qualcomm QNN / MediaPipe"],
        ["Speed", "Interactive Baseline", "Target: ≥24 tok/s"]
    ]

    for r_idx, row in enumerate(matrix_data):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if r_idx == 0 else RGBColor(12, 15, 24)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(10)
            p.font.bold = (r_idx == 0 or c_idx == 0)
            if r_idx == 0:
                p.font.color.rgb = IQOO_AMBER if c_idx == 1 else WHITE
            elif c_idx == 1:
                p.font.color.rgb = IQOO_GREEN
            else:
                p.font.color.rgb = WHITE

    # ==========================================
    # SLIDE 5: LIVE DEMO & IMPACT
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "The 60-Second Live Stage Demo & Impact", "04 / DEMO & CONCLUSION")

    c_demo = add_card(s5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.9))
    tfd = c_demo.text_frame
    tfd.word_wrap = True
    pt = tfd.paragraphs[0]
    pt.text = "🎬 The 60-Second Airplane Mode Demo"
    pt.font.size = Pt(16)
    pt.font.bold = True
    pt.font.color.rgb = IQOO_AMBER

    demo_flow = (
        "1. 0:00 - 0:15 | The Physical Proof:\n"
        "   Swipe notification shade, toggle Airplane Mode.\n"
        "   Status badge reacts instantly: 'AIRPLANE ON • 0 KB'.\n\n"
        "2. 0:15 - 0:35 | The Scan & Retrieval:\n"
        "   Point camera at textbook incline physics problem.\n"
        "   Offline OCR extracts question; local RAG finds concept.\n\n"
        "3. 0:35 - 0:50 | Socratic Dialogue:\n"
        "   EdgeTutor asks: 'What forces act along the slope?'\n"
        "   Student guided step-by-step with KaTeX formulas.\n\n"
        "4. 0:50 - 1:00 | Verifiable 0 KB Result:\n"
        "   Show system network stats: Exactly 0 KB consumed."
    )
    pd = tfd.add_paragraph()
    pd.text = demo_flow
    pd.font.size = Pt(11)
    pd.font.color.rgb = WHITE
    pd.space_before = Pt(8)

    c_impact = add_card(s5, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.9))
    tfi = c_impact.text_frame
    tfi.word_wrap = True
    ptr = tfi.paragraphs[0]
    ptr.text = "🚀 Strategic Impact & Brand Value"
    ptr.font.size = Pt(16)
    ptr.font.bold = True
    ptr.font.color.rgb = IQOO_CYAN

    impact_txt = (
        "• Redefines Performance Beyond Gaming:\n"
        "  Proves iQOO silicon isn't just for BGMI—it's India's smartest productivity engine.\n\n"
        "• The 'Parent-Approved' Buying Decision:\n"
        "  Parents gladly invest in an iQOO phone when it serves as a distraction-free offline study mentor.\n\n"
        "• Potential Out-of-the-Box Feature:\n"
        "  Strong candidate for pre-loading as a flagship educational capability on FuntouchOS.\n\n"
        "🌟 Final Punchline:\n"
        "  'Don't distract students with the internet.\n"
        "   Empower them with On-Device AI.'"
    )
    pdi = tfi.add_paragraph()
    pdi.text = impact_txt
    pdi.font.size = Pt(11)
    pdi.font.color.rgb = WHITE
    pdi.space_before = Pt(8)

    output_path = r"c:\Users\varanasi saivinay\OneDrive\private\IQ\submission\iQOO_EdgeTutor_PitchDeck.pptx"
    prs.save(output_path)
    print(f"Master pitch deck saved to: {output_path}")

    downloads_path = os.path.expanduser(r"~\Downloads\iQOO_EdgeTutor_PitchDeck.pptx")
    try:
        shutil.copy2(output_path, downloads_path)
        print(f"Copied master pitch deck to: {downloads_path}")
    except Exception as e:
        print(f"Could not copy to Downloads: {e}")

if __name__ == "__main__":
    create_deck()
